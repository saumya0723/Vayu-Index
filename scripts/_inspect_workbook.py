"""
One-off workbook inspector — does NOT modify the source file.
Run: python scripts/_inspect_workbook.py
"""
import os
import hashlib
import openpyxl
import pandas as pd

RAW_PATH = os.path.join("data", "official", "raw", "cpi_1822(final).xlsx")

# ---- SHA-256 hash (proof of immutability) -----------------------------------
with open(RAW_PATH, "rb") as f:
    raw_bytes = f.read()
sha256 = hashlib.sha256(raw_bytes).hexdigest()
print("File       :", RAW_PATH)
print("Size       :", os.path.getsize(RAW_PATH), "bytes")
print("SHA-256    :", sha256)
print()

# ---- Sheet names -----------------------------------------------------------
wb = openpyxl.load_workbook(RAW_PATH, read_only=True, data_only=True)
print("Sheet names:", wb.sheetnames)
print()

# ---- Per-sheet deep inspection (openpyxl raw + pandas) --------------------
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print("=" * 60)
    print("SHEET:", sheet_name)
    print("=" * 60)

    all_rows = list(ws.iter_rows(values_only=True))
    print("Total rows (incl. header):", len(all_rows))
    print()

    print("RAW rows 0-9:")
    for i, row in enumerate(all_rows[:10]):
        print("  [%02d]" % i, row)
    print()

    if len(all_rows) > 10:
        print("RAW last 5 rows:")
        for i, row in enumerate(all_rows[-5:], len(all_rows) - 5):
            print("  [%02d]" % i, row)
        print()

wb.close()

# ---- Also read with pandas (header=0) to check column alignment -----------
print("=" * 60)
print("PANDAS read (header=0):")
print("=" * 60)
try:
    df = pd.read_excel(RAW_PATH, sheet_name=None, header=0)
    for sname, sdf in df.items():
        print("Sheet:", sname)
        print("  Shape:", sdf.shape)
        print("  Columns:", list(sdf.columns))
        print("  Dtypes:")
        print(sdf.dtypes.to_string())
        print()
        print("  Head(10):")
        print(sdf.head(10).to_string())
        print()
        print("  Tail(5):")
        print(sdf.tail(5).to_string())
        print()
        print("  Missing values per column:")
        print(sdf.isnull().sum().to_string())
        print()
        print("  Unique values per column (first 20 uniques each):")
        for col in sdf.columns:
            uvals = sdf[col].dropna().unique()
            print("    %s: %s" % (col, uvals[:20]))
        print()
except Exception as e:
    print("pandas read error:", e)
